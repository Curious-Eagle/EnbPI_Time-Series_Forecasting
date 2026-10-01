"""Build the beginner guide using bundled python-docx and native Word equations."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'Understanding_the_EnbPI_Project.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.78)
sec.footer_distance = Inches(.28)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 26), ('Subtitle', 13), ('Heading 1', 19), ('Heading 2', 13)]:
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.color.rgb = 'Calibri', Pt(size), RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(9 if name == 'Heading 2' else 0)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
    pr = st.element.find(qn('w:pPr'))
    if pr is not None:
        for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)
for name in ['List Bullet', 'List Number']:
    doc.styles[name].font.name, doc.styles[name].font.size = 'Calibri', Pt(11)
    doc.styles[name].paragraph_format.space_after = Pt(5)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('EnbPI project guide  |  ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Understanding the EnbPI Project'
doc.core_properties.subject = 'A beginner guide to the implemented research reproduction'
doc.core_properties.author = 'EnbPI portfolio project'
doc.core_properties.keywords = 'EnbPI, forecasting, prediction intervals, beginner guide'


def p(text, style=None):
    para = doc.add_paragraph(text, style)
    para.paragraph_format.widow_control = True
    return para


def h(text): return doc.add_heading(text, level=2)


def page(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)


def bullet(text): return p(text, 'List Bullet')


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, label in enumerate(headers): t.rows[0].cells[i].text = label
    for values in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, values): cell.text = str(value)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6EF' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:' + side)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: border.set(qn('w:' + key), value)
                borders.append(border)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '65'), ('bottom', '65'), ('left', '95'), ('right', '95')]:
                child = OxmlElement('w:' + side); child.set(qn('w:w'), value); child.set(qn('w:type'), 'dxa'); margins.append(child)
            tcpr.append(margins)
            for para in cell.paragraphs:
                if headers[ci] in {'Count', 'Target', 'Local coverage', 'Author coverage', 'Local width', 'Author width'}:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.02
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.bold = ri == 0
    p('').paragraph_format.space_after = Pt(1)
    return t


def mr(text):
    run = OxmlElement('m:r')
    el = OxmlElement('m:t'); el.text = str(text); run.append(el)
    return run


def append_math(parent, parts):
    for part in parts:
        parent.append(mr(part) if isinstance(part, (str, int, float)) else part)


def sub(base, index):
    el = OxmlElement('m:sSub'); e = OxmlElement('m:e'); s = OxmlElement('m:sub')
    append_math(e, [base]); append_math(s, [index]); el.extend([e, s]); return el


def sup(base, power):
    el = OxmlElement('m:sSup'); e = OxmlElement('m:e'); s = OxmlElement('m:sup')
    append_math(e, [base]); append_math(s, [power]); el.extend([e, s]); return el


def frac(numerator, denominator):
    el = OxmlElement('m:f'); n = OxmlElement('m:num'); d = OxmlElement('m:den')
    append_math(n, numerator if isinstance(numerator, list) else [numerator])
    append_math(d, denominator if isinstance(denominator, list) else [denominator])
    el.extend([n, d]); return el


def sigma(low, high, expression):
    el = OxmlElement('m:nary'); props = OxmlElement('m:naryPr')
    char = OxmlElement('m:chr'); char.set(qn('m:val'), '∑'); props.append(char)
    loc = OxmlElement('m:limLoc'); loc.set(qn('m:val'), 'undOvr'); props.append(loc)
    el.append(props)
    for tag, content in [('sub', [low]), ('sup', [high]), ('e', expression)]:
        child = OxmlElement('m:' + tag); append_math(child, content); el.append(child)
    return el


def eq(*parts):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.keep_together = True
    mathpara = OxmlElement('m:oMathPara'); math = OxmlElement('m:oMath')
    append_math(math, parts); mathpara.append(math); para._p.append(mathpara)
    return para


def picture(path, caption, width=6.6):
    para = doc.add_paragraph()
    para.paragraph_format.keep_with_next = True
    run = para.add_run(); shape = run.add_picture(str(ROOT / path), width=Inches(width))
    shape._inline.docPr.set('descr', caption)
    cap = p(caption); cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs: run.font.size = Pt(9)


def code(lines):
    for line in lines:
        para = p(line)
        para.paragraph_format.space_after = Pt(2)
        for run in para.runs: run.font.name, run.font.size = 'Consolas', Pt(9)


def link(label, url):
    para = doc.add_paragraph()
    relationship = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    el = OxmlElement('w:hyperlink'); el.set(qn('r:id'), relationship)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '174B70'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text)
    el.append(run); para._p.append(el)


# Page 1
doc.add_paragraph('Understanding the EnbPI Project', 'Title')
doc.add_paragraph('A practical guide from the first idea to the saved results', 'Subtitle')
p('This project teaches a model to predict a future number and provide a useful range around that prediction. You will learn what the model does, how EnbPI builds the range, how the code works, and what the results actually tell us.')
p('You do not need advanced mathematics to begin. Read the explanation before each formula, then work through the small examples. Every mathematical symbol is introduced before you need it.')
h('The project in one paragraph')
p('We reproduced one part of a 2021 research paper by Chen Xu and Yao Xie. We used past hourly solar measurements to predict the next measurement, then built uncertainty intervals with EnbPI. At a 90% coverage target, the intervals covered about 90.24% of test observations. The authors\' saved result was about 90.25%. The next application is retail sales forecasting.')
h('What is complete')
bullet('A partial reproduction on the authors\' solar dataset: 10 random seeds and five coverage targets.')
bullet('Code checks, comparisons with the original implementation, saved predictions, charts, and reports.')
bullet('An investigation of coverage by hour, including weaker coverage around midday.')
h('What comes later')
p('Retail sales data, inventory experiments, and a retail dashboard are planned. They have not been evaluated yet. The project folder is named for the eventual retail application, but the completed results are from solar data.')
h('How to read this guide')
table(['Read', 'What you will understand'], [
('Pages 2 to 5', 'Forecasting, the dataset, trees, and bootstrap sampling'),
('Pages 6 to 8', 'The complete EnbPI process and worked formulas'),
('Pages 9 to 12', 'Evaluation, actual results, and the limits of the method'),
('Pages 13 to 17', 'Code, commands, retail plans, interview answers, and references')], [1.5, 5.15])
p('Keep the project open while reading. The walkthrough notebook is useful after you finish the explanation of the algorithm.')

# Page 2
page('1 The problem we are solving')
p('Imagine a shop that expects to sell 100 units tomorrow. Ordering exactly 100 may be risky: tomorrow could be unusually busy or quiet. A forecast range adds information about how uncertain the prediction is. This is a motivating example, not a result from our current experiment.')
h('A forecast and an interval answer different questions')
table(['Output', 'Example', 'Meaning'], [
('Point forecast', '100 units', 'A single predicted value'),
('Prediction interval', '80 to 130 units', 'A range intended to contain the future observation'),
('Nominal coverage', '90%', 'The coverage frequency we aim for')], [1.35, 1.3, 4.0])
p('A 90% target means we want the procedure to cover future outcomes about 90% of the time under suitable conditions. It does not promise coverage on every date, and it is not a guarantee that a particular item tomorrow has exactly a 90% chance of falling inside its interval.')
h('Why uncertainty exists')
p('Real outcomes vary. Weather changes solar irradiance; promotions and customer behavior change sales. A model also makes mistakes because it has limited data and cannot capture every pattern. EnbPI uses past prediction errors to help choose the size of a range. It does not explain the cause of every error.')
h('The basic notation')
table(['Symbol', 'Read it as', 'Meaning'], [
('t', 'time t', 'The hour or day being predicted'),
('yₜ', 'y at time t', 'The actual value'),
('cₜ', 'center at time t', 'The center of our prediction interval'),
('Lₜ and Uₜ', 'lower and upper', 'The two ends of the interval'),
('α', 'alpha', 'The target miss rate; 0.10 means 10%')], [1.15, 1.65, 3.85])
eq('Target coverage = 1 − α = 1 − 0.10 = 0.90')
p('A prediction interval concerns a future observation. A confidence interval usually concerns an unknown parameter, such as an average. In this project, we are producing prediction intervals.')

# Page 3
page('2 The data and the time order')
p('The completed experiment uses 8,760 hourly observations from Atlanta in 2018. The target is diffuse horizontal irradiance, abbreviated DHI, measured in watts per square metre. It describes part of the sunlight reaching a horizontal surface. We predict its next hourly value.')
table(['Part of the series', 'Count', 'Purpose'], [
('Initial training observations', '1,752', 'First 20% of the year'),
('Usable training examples', '1,732', 'Training observations after building 20 lags'),
('Test observations', '7,008', 'Later values used for evaluation')], [2.6, .85, 3.2])
h('A lag is simply a past value')
p('The inputs for time t are the previous 20 measurements. The target is the measurement at t. The oldest input comes first in the code. Weather columns exist in the source file, but this selected experiment uses only past DHI values.')
eq(sub('X', 't'), ' = [', sub('y', 't−20'), ', …, ', sub('y', 't−1'), ']')
p('X means the input vector, or list of features. The number y at time t is the answer we want the model to predict. A simpler three-lag example makes the arrangement clear:')
table(['Past values used as inputs', 'Next value used as target'], [
('10, 12, 15', '14'), ('12, 15, 14', '18'), ('15, 14, 18', '17')], [3.4, 3.25])
p('The first 20 training observations supply history but cannot each form a full 20-lag example. That is why 1,752 − 20 = 1,732 usable training examples remain.')
h('Why we do not shuffle time')
p('Training on later observations and testing on earlier ones would not represent real forecasting. We train on the earlier portion and test on the later portion. This also avoids using the answer to create its own inputs, a mistake called data leakage.')
p('During the test period we predict one hour, observe its actual value, and then move to the next hour. Previous test outcomes can therefore become legitimate lag inputs. This is a rolling one-hour-ahead evaluation, not a forecast of the whole year made in advance.')
p('The first test target is 15 March 2018 at 00:30. The final target is 31 December 2018 at 23:30. Timestamps retain the dataset\'s local standard time.')

# Page 4
page('3 How a decision tree makes a prediction')
p('A regression model predicts a number. A decision tree does this by repeatedly dividing training examples into groups using questions about their inputs. A possible question is whether the previous hour\'s irradiance was above a threshold.')
h('A small tree example')
p('Suppose one group contains target values 20, 30, and 40. Under squared-error training, the prediction in that final group, called a leaf, is their average:')
eq('Leaf prediction = ', frac('20 + 30 + 40', '3'), ' = 30')
p('The tree chooses splits that reduce squared prediction errors across the resulting groups. If a prediction misses the actual value by 10, its squared error is 100. If it misses by 20, its squared error is 400, so large mistakes receive a larger penalty.')
eq('Squared error = ', sup('(actual − prediction)', '2'))
p('A tree with maximum depth 2 can make at most two splits along any route from the root to a leaf. That allows at most four leaves. It is a deliberately small model for this reproduction, not a claim that depth 2 is best for every forecasting task.')
h('How a forest combines trees')
p('A random-forest regressor combines several regression trees and averages their predictions. For example, tree predictions of 90, 100, and 110 average to 100. Our base regressor has 10 trees. [3]')
eq('Forest prediction = ', frac('Sum of the 10 tree predictions', '10'))
h('The exact settings in this project')
table(['Setting', 'Value', 'Meaning'], [
('n_estimators', '10', 'Ten trees in each forest'),
('max_depth', '2', 'At most two splits along a route'),
('criterion', 'squared_error', 'Choose splits using squared errors'),
('bootstrap', 'False', 'No extra row resampling inside each forest'),
('max_features', '1.0', 'Consider all 20 features at each split')], [1.65, 1.1, 3.9])
p('The main diversity here comes from EnbPI\'s outer bootstrap samples. Because the inner forest sees the same rows and all features, its trees can be very similar or identical, apart from tie-related randomness. Do not describe this setup as using every default random-forest setting. [3]')

# Page 5
page('4 Bootstrap sampling and out of bag errors')
p('Bootstrap sampling means drawing training rows with replacement. Replacement means that after a row is chosen, it can be chosen again. A sample has the same number of rows as the original training set, but some rows repeat and others are absent.')
h('A tiny example')
table(['Bootstrap sample', 'Selected row names', 'Rows left out'], [
('A', '1, 1, 2, 4, 5', '3'),
('B', '2, 2, 3, 4, 4', '1 and 5'),
('C', '1, 3, 3, 5, 5', '2 and 4')], [1.5, 3.15, 2.0])
p('We train a different forest on each bootstrap sample. To estimate an error for training row 3, we use only forests whose samples left row 3 out. In the example, forest A qualifies; forests B and C do not.')
p('A row left out of a bootstrap sample is called out of bag, or OOB, for that model. This helps avoid measuring an error using a model that directly trained on the very same row. It does not make neighboring time-series observations independent.')
h('Turn the prediction into an error')
p('Suppose the actual value for a training row is 120 and the average prediction from its OOB models is 105. Its absolute residual is 15. Absolute means we ignore whether the model was too high or too low.')
eq('Absolute residual = |120 − 105| = 15')
p('Repeating this for all 1,732 training examples produces the initial error collection used to size intervals. EnbPI therefore does not need a separate calibration split for this selected setup.')
h('Why 30 bootstrap samples are useful')
p('With n draws from n rows, the chance a particular row is never selected is:')
eq('P(left out) = ', sup('(1 − 1/n)', 'n'), ' ≈ 0.368')
p('For a large training set, a row is left out of about 36.8% of bootstrap samples on average. With 30 bootstrap forests, that is roughly 11 eligible forests per row. The actual number varies. If none qualify, the original implementation uses a zero prediction; our reproduction retains and records that behavior.')

# Page 6
page('5 The complete EnbPI process')
p('EnbPI stands for Ensemble Batch Prediction Intervals. Ensemble means that models are combined. Batch describes the possibility of issuing several intervals before receiving new outcomes. Our experiment uses a batch size of one, so feedback arrives after each hourly prediction. [1, 2]')
h('Training happens once per seed')
for text in [
    'Build lag inputs from the initial training period.',
    'Draw 30 bootstrap samples and fit one 10-tree forest to each sample.',
    'For every training row, average predictions from forests that omitted that row.',
    'Calculate the absolute training residuals and keep them as the initial calibration window.'
]: p(text, 'List Number')
h('Prediction repeats for each test hour')
for text in [
    'Build the input from the latest 20 observed measurements.',
    'Use the stored forests to calculate the interval center. The exact rule is explained on the next page.',
    'Find the selected percentile of the current error window. This is the interval radius.',
    'Issue the lower and upper bounds before observing the current target.',
    'When the actual value arrives, calculate its absolute error, remove the oldest error, and append the new error.'
]: p(text, 'List Number')
eq(sub('L', 't'), ' = ', sub('c', 't'), ' − ', sub('q', 't'), '     ', sub('U', 't'), ' = ', sub('c', 't'), ' + ', sub('q', 't'))
p('Here q is the radius. The full width is 2q. For a center of 100 and a radius of 20, the interval is [80, 120] and the width is 40.')
h('What changes and what stays fixed')
p('The fitted forests stay fixed during testing. The interval radius changes as the error window changes. The input lags also change because new actual measurements become available. Updating an error window is not the same as retraining the model.')
p('There are 30 forests × 10 trees = 300 fitted trees per seed. We reuse that ensemble at five alpha values. Across ten seeds, there are 3,000 tree fits and 50 evaluated seed-and-alpha combinations.')

# Page 7
page('6 The exact prediction center')
p('This page explains a detail that is easy to miss. The conference implementation does not simply average all 30 forest predictions to form an interval center. It constructs a prediction for each leave-one-out group and then selects an upper quantile of those predictions. [1, 2]')
h('First define the leave one out prediction')
p('Let f with subscript b mean the prediction from bootstrap forest b. Let O with subscript i be the set of forests that omitted training row i. Let g with subscript i be their average prediction for any input X.')
p('The symbol ∑ means add up. The expression b ∈ Oᵢ means use only the forests belonging to the group Oᵢ.')
eq(sub('g', 'i'), '(X) = ', frac([sigma('b∈Oᵢ', '', [sub('f', 'b'), '(X)'])], '|Oᵢ|'))
p('The vertical bars around O mean the number of qualifying forests. For example, if three qualifying forests predict 90, 100, and 110, their average is 100. For the training row itself, the initial error is:')
eq(sub('e', 'i'), ' = |', sub('y', 'i'), ' − ', sub('g', 'i'), '(', sub('X', 'i'), ')|')
h('Then choose a center for a new input')
p('For the new input at time t, calculate one g prediction for each of the n training rows. Sort those n predictions from smallest to largest. The code chooses the item at the following zero-based index:')
eq('k = ⌊(1 − α)n⌋')
p('The floor brackets mean round down to an integer. Zero-based index 0 is the first item. With n = 1,732 and alpha = 0.10, k = floor(1,558.8) = 1,558, so the selected value is the 1,559th item in the sorted list. That selected value is the center c at time t.')
h('A miniature center example')
p('Suppose the five sorted group predictions are 90, 95, 100, 105, and 110. If alpha is 0.20, the index is floor(0.8 × 5) = 4. The selected center is 110. This illustrates the exact order-statistic convention, using deliberately tiny invented data.')
p('The center can change with alpha even when the fitted forests stay the same. Keep this behavior when reproducing the conference code. Replacing it with an ordinary average would create a different experiment.')

# Page 8
page('7 Build an interval with actual numbers')
p('We now have a center. To determine the radius, we use a percentile of recent absolute errors. A percentile tells us where a value lies in an ordered list. The 90th percentile is near the high end: approximately 90% of the reference values lie at or below that part of the distribution.')
h('A complete small example')
p('Use these invented residuals: [2, 4, 6, 8, 10]. Suppose the next forecast center is 100 and alpha is 0.10. Our NumPy calculation uses linear interpolation. For m sorted values and percentile proportion p, the position is: [4]')
eq('r = (m − 1)p = (5 − 1) × 0.90 = 3.6')
p('Positions start at zero. Position 3 holds 8; position 4 holds 10. Position 3.6 is 60% of the way between them.')
eq('q = 8 + 0.6 × (10 − 8) = 9.2')
eq('[L, U] = [100 − 9.2, 100 + 9.2] = [90.8, 109.2]')
p('The interval width is 109.2 − 90.8 = 18.4. Now reveal the actual value, 112. It is outside the interval, and its new absolute error is |112 − 100| = 12.')
h('Update only after the observation arrives')
table(['Stage', 'Error window'], [
('Before observing 112', '2, 4, 6, 8, 10'),
('Remove oldest error and append 12', '4, 6, 8, 10, 12')], [3.5, 3.15])
p('The updated 90th percentile is 10 + 0.6 × (12 − 10) = 11.2. If the next center were again 100, its interval would be [88.8, 111.2]. It widened because a larger error entered the window.')
p('In the real experiment the window contains 1,732 errors. It is ordered by when errors become available; percentile calculation sorts a copy as needed. We remove the oldest error, not the smallest one. Width can also shrink as large old errors leave.')
h('Two quantile rules are intentionally different')
p('The center uses the sorted-item rule on page 7. The radius uses NumPy\'s interpolated percentile, as shown above. The original code first computes int(100 × (1 − alpha)) for that percentile. We preserve that convention instead of adding a different finite-sample correction.')

# Page 9
page('8 How we judge a prediction interval')
p('We need both coverage and width. A range from minus one million to plus one million might cover almost everything but would rarely help a decision. We also measure the error in the center and penalize missed intervals.')
h('Coverage')
eq('Coverage = ', frac('Number of actual values inside their intervals', 'Number of evaluated predictions'))
p('Count an outcome as covered when lower ≤ actual ≤ upper, including the boundaries. If 90 of 100 actuals fall inside their intervals, coverage is 0.90, or 90%. This is an interval metric, not classification accuracy.')
h('Average width')
eq('Mean width = ', frac([sigma('t=1', 'N', ['(', sub('U', 't'), ' − ', sub('L', 't'), ')'])], 'N'))
p('N is the number of test observations. If three widths are 10, 20, and 30, their average is 20. Smaller widths are useful only when coverage is also acceptable.')
h('Mean absolute error')
eq('MAE = ', frac([sigma('t=1', 'N', ['|', sub('y', 't'), ' − ', sub('c', 't'), '|'])], 'N'))
p('If the center errors are 2, 4, and 6, MAE is 4. This evaluates the center prediction, not the interval. In this implementation, MAE can vary slightly with alpha because the center depends on alpha.')
h('Interval score')
eq('IS = (U − L) + ', frac('2', 'α'), 'max(L − y, 0) + ', frac('2', 'α'), 'max(y − U, 0)')
p('Here max(a, 0) means use a if it is positive; otherwise use zero. If the actual is inside the interval, both penalties are zero and the score equals the width. If it is outside, the score adds a penalty for the distance missed. Average this score across observations. Lower is better when comparing the same alpha and target units.')
p('Example: interval [80, 120], alpha 0.10. If actual = 100, score = 40. If actual = 130, score = 40 + (2/0.10) × 10 = 240. The much larger score reflects the missed outcome.')

# Page 10
page('9 The experiment we actually ran')
p('The experiment reproduces the random-forest, lag-only portion of Figure 1 in the ICML 2021 paper. It does not reproduce every model or every figure. The reference numbers come from the authors\' saved CSV file, rather than reading values from a chart.')
table(['Setting', 'Value'], [
('Data', 'Atlanta solar DHI, 2018, 8,760 hourly rows'),
('Initial training portion', '20%, followed by a chronological test period'),
('Inputs', 'Previous 20 target observations'),
('Ensemble', '30 bootstrap forests, 10 trees each, depth 2'),
('Feedback', 'One observed outcome after each forecast'),
('Target coverage levels', '75%, 80%, 85%, 90%, 95%'),
('Random seeds', '98765 through 98774'),
('Total evaluated settings', '10 seeds × 5 levels = 50')], [2.3, 4.35])
h('Why repeat a random seed')
p('A seed controls random sampling. Repeating the same seed and software environment should reproduce the same sampling sequence. Different seeds show how sensitive results are to bootstrap randomness. They all use the same observed year; they are not ten independent datasets.')
h('Local results against the reference')
table(['Target', 'Local coverage', 'Author coverage', 'Local width', 'Author width'], [
('75%', '75.15%', '75.15%', '95.63', '95.62'),
('80%', '79.96%', '79.93%', '125.37', '125.47'),
('85%', '85.15%', '85.16%', '169.87', '169.97'),
('90%', '90.24%', '90.25%', '256.64', '256.66'),
('95%', '95.11%', '95.11%', '373.90', '373.88')], [.75, 1.5, 1.5, 1.45, 1.45])
p('Widths are in W/m². These values are averages across ten seeds and are rounded for reading. The saved CSV files contain the full precision.')
p('Close agreement supports this selected reproduction. It does not establish that EnbPI is best for every dataset or that the retail application will perform equally well.')

# Page 11
page('10 What the results tell us')
picture('outputs/solar_reproduction/coverage_width.png', 'The local results closely follow the saved author results for the selected experiment.')
p('As the coverage target rises, average interval width increases. At the 90% target, the average full width is about 256.64 W/m². This is a full width, not a plus-or-minus radius of 256.64.')
picture('outputs/solar_reproduction/hourly_coverage.png', 'Coverage by local hour at the 90% target, averaged across the ten seeds.')
p('Overall coverage is about 90.24%, but coverage at hour 12 is only about 61.20%. Several nighttime hours have 100% coverage. Each hour has 292 test observations per seed. This difference is why we inspect subgroups as well as the overall average.')
p('Nighttime zeros are easier to include inside intervals. Greater daytime variation is a plausible contributor to the midday pattern, but this experiment does not establish a single cause for every miss.')
h('The frozen error comparison')
p('As an extra diagnostic, we kept the same centers but stopped updating the initial error distribution. At the 90% target, this gave about 79.55% coverage and width 117.34 W/m². Its narrower intervals missed more outcomes. This comparison illustrates the effect of updating; it is not an original Figure 1 baseline.')

# Page 12
page('11 What the method can and cannot guarantee')
h('Overall coverage and subgroup coverage')
p('Marginal coverage describes coverage averaged over the relevant distribution. Conditional coverage asks about a particular input or group, such as midday observations. A method can perform well on the average and poorly in one group. Our hourly results show exactly why the distinction matters.')
h('What conformal means here')
p('Conformal methods build prediction sets using scores that describe how unusual a new outcome would be relative to reference observations. Here the score is an absolute prediction error. EnbPI is related to this framework and uses bootstrap predictions plus sequential error updates for time series. [1, 2]')
p('Classic conformal arguments often assume exchangeability: roughly, the joint data behavior is unchanged by reordering observations. Time series generally do not have that property. EnbPI\'s theory instead provides approximate coverage under conditions on temporal errors and model quality. [1]')
h('The assumptions in ordinary language')
bullet('The errors must have suitable stability and dependence properties. In the strongly mixing case, sufficiently distant errors become weakly dependent.')
bullet('The fitted predictors must approximate the relevant underlying regression function well enough for the theoretical bounds to be useful.')
bullet('The feedback schedule and available inputs must match the forecasting setup.')
p('Distribution-free does not mean assumption-free. The method does not require choosing a specific error distribution such as a normal distribution, but it does not promise success for arbitrary data or a completely unsuitable model. Our experiment measures empirical performance; it does not prove those assumptions hold.')
h('Practical limitations of this reproduction')
bullet('It covers one location, one year, one target, and one model family. Later years and retail products are untested.')
bullet('Intervals may include negative solar values. We kept the original bounds for comparability.')
bullet('The fitted models do not retrain during the test period. Sudden changes can make their predictions less useful.')
bullet('Increasing the target coverage usually widens intervals, but it does not guarantee the target is achieved in every window.')
p('A seed standard deviation measures variation across randomized fits on the same observations. It is not a confidence interval for performance in a new year.')

# Page 13
page('12 How the ideas map to the code')
p('All paths below are relative to the enbpi-demand-forecasting project folder. Read the files in this order so that inputs, predictions, and evaluation form one connected story.')
table(['File or function', 'What it does'], [
('configs/reproduction.json', 'Stores the experiment settings, seeds, and alpha values.'),
('src/enbpi_project/data.py', 'Loads the original data and builds chronological lag inputs.'),
('EnbPI.fit in model.py', 'Draws bootstrap samples, fits forests, and computes initial OOB errors.'),
('EnbPI.predict_centers', 'Calculates the alpha-dependent centers without receiving future targets.'),
('RollingIntervals.interval', 'Computes the radius and returns the bounds from the current errors.'),
('RollingIntervals.observe', 'Adds the error only after the actual outcome becomes available.'),
('sequential_intervals', 'Calls interval first and observe second for every test outcome.'),
('src/enbpi_project/metrics.py', 'Calculates coverage, width, MAE, and interval score.'),
('scripts/run_experiment.py', 'Connects the steps, saves predictions, compares results, and draws charts.')], [2.55, 4.1])
h('A simple version of the prediction loop')
code(['For each test time t:', '    use past observations to build the input', '    calculate the forecast center', '    calculate the interval from the current error window', '    save the interval', '    observe the actual value', '    remove the oldest error and append the new error'])
p('The code can prepare the lag matrix in advance because this is an offline replay of one-step forecasting. Each row is checked to contain only values from earlier times. It does not mean every target was available to a live forecaster at the start.')
p('A deque is the fixed-length queue used for errors. NumPy handles numerical arrays, pandas handles tables, scikit-learn fits forests, matplotlib draws charts, and pytest runs validation checks.')

# Page 14
page('13 Run the project and inspect the outputs')
p('Open PowerShell in the project folder. The existing .venv directory contains the isolated Python environment. The following commands use it directly, so you do not need to activate it first.')
h('Run the checks')
code([r'.\.venv\Scripts\python.exe -m pytest'])
p('The 11 checks include numerical agreement with selected unchanged author methods at five alpha values, correct lag construction, preventing current or future outcomes from changing issued intervals, and known-example metric calculations.')
h('Run one small experiment')
code([r'.\.venv\Scripts\python.exe scripts/run_experiment.py --config configs/smoke.json'])
p('This is one command. It uses one seed and 90% nominal coverage and saves results under outputs/solar_smoke.')
h('Run the selected full reproduction')
code([r'.\.venv\Scripts\python.exe scripts/run_experiment.py --config configs/reproduction.json'])
p('This evaluates ten seeds and five alpha values. Re-running a configuration replaces its generated output files. Training does not need a GPU or a paid API.')
h('Open the useful results first')
table(['Inside outputs/solar_reproduction', 'How to use it'], [
('report.md', 'Read the findings and limitations.'),
('summary.csv', 'Compare means and seed variation.'),
('upstream_comparison.csv', 'Inspect differences from author results.'),
('predictions/*.csv.gz', 'Inspect actuals, centers, bounds, and hits.'),
('hourly_coverage.csv', 'Find hours where coverage is weak.'),
('config.json and run_metadata.json', 'Check settings, source version, packages, and dates.')], [3.3, 3.35])
p('The walkthrough notebook, notebooks/01_explore_reproduction.ipynb, reads saved results. It does not train models. It needs a Jupyter-capable editor and a suitable Python kernel; the command-line experiments do not.')
p('Original downloaded files are preserved under references/upstream. The source manifest records their URLs and SHA-256 checksums, which detect changes to file contents. A checksum supports traceability; it does not prove scientific correctness.')

# Page 15
page('14 How this becomes a retail project')
p('The completed solar experiment establishes that we can reproduce a selected published result. The retail extension asks a new question: can the same approach give useful next-day sales intervals for different products? Success on solar data does not answer that question by itself.')
h('Translate the ingredients')
table(['Solar reproduction', 'Proposed retail extension'], [
('Hourly DHI', 'Daily units sold for a product'),
('Previous 20 hourly values', 'Past daily sales and suitable calendar features'),
('One-hour-ahead target', 'Next-day target'),
('Feedback after each hour', 'Feedback after each day'),
('Coverage by hour', 'Coverage by product and sales pattern')], [3.0, 3.65])
h('A fair experiment')
p('Select a manageable product subset using training-period information only. Use an early period for training, a later development period for choosing settings, and a final untouched period for evaluation. This prevents repeated test-based adjustments from making results look better than they really are.')
p('Compare EnbPI with a simple seasonal-naive point forecast, such as using sales from the same weekday last week, and a clearly defined residual-interval baseline. Measure coverage, width, interval score, and point error on the same dates and information schedule.')
h('Features must be available when the forecast is made')
p('Tomorrow\'s weekday is already known. Tomorrow\'s observed sales are not. A promotion may be usable if it was scheduled in advance. Any rolling sales average must exclude the day being predicted. These details matter more than adding an elaborate model.')
h('Sales are not always true demand')
p('If a shop has only five units in stock and sells all five, customer demand may have been higher. The sales record alone cannot reveal all lost demand. We should call the observed target sales and state this limitation unless inventory information is available.')
p('A prediction interval also does not directly specify an optimal order quantity. That requires assumptions about lead times, shortages, holding costs, and replenishment policy. Any inventory-cost simulation would need those assumptions stated explicitly.')
p('A useful original contribution would test whether apparently good overall coverage hides poor intervals for particular product groups. This follows naturally from the hourly weakness found in the solar reproduction.')

# Page 16
page('15 Explain the work in an interview')
h('A short explanation you can make your own')
p('I reproduced the lag-only random-forest portion of an ICML 2021 paper on prediction intervals for time series. I used EnbPI to combine bootstrap models with a rolling window of prediction errors. Across ten seeds, the 90% target gave about 90.24% coverage, close to the authors\' saved result. I also found much weaker midday coverage, showing why overall calibration alone is not enough. The retail application is the next stage.')
h('Questions you should be ready to answer')
table(['Question', 'A clear answer'], [
('Why EnbPI?', 'It builds sequential uncertainty intervals around fitted models and can use dependent time-series observations under stated assumptions.'),
('Why bootstrap?', 'It creates different training samples and leaves rows out, allowing errors to be estimated from models that omitted those rows.'),
('Did you train on the test set?', 'The forests were fitted only on the initial training period. Earlier test actuals were used later as lag inputs and feedback, as required by one-step forecasting.'),
('Why not report only MAE?', 'MAE measures center error. It does not tell us whether prediction intervals cover enough outcomes or are too wide.'),
('What did you contribute?', 'A focused runner, compatibility documentation, validation checks, result comparisons, and hourly and rolling diagnostics. The core method comes from the paper.'),
('What is still unproven?', 'Performance on retail products, later years, different horizons, and inventory decisions.')], [2.15, 4.5])
h('Check your own understanding')
p('1. A center is 50 and the radius is 8. What are the bounds and full width?\n2. Why can the newest actual value update the next interval but not its own?\n3. Does 90% overall coverage guarantee 90% coverage at midday?\n4. If old large errors leave the window, can the interval become narrower?')
p('Answers: 1. [42, 58], width 16. 2. It arrives after its own interval is issued. 3. No; subgroup coverage can differ. 4. Yes, if the chosen percentile decreases.')
p('Use the code and the numeric example to explain your reasoning rather than memorizing a script. Credit the paper and reused code, and be transparent about coding-assistant support.')

# Page 17
page('16 A quick reference for later')
table(['Term', 'Simple meaning'], [
('Time series', 'Measurements ordered through time.'),
('Feature', 'An input used to make a prediction.'),
('Lag', 'A past target value used as an input.'),
('Regression', 'Predicting a numerical value.'),
('Bootstrap', 'Sampling rows with replacement.'),
('Out of bag', 'A row was not selected for a particular model.'),
('Ensemble', 'A collection of models whose predictions are combined.'),
('Residual', 'An observed value minus its prediction; this project uses its absolute size.'),
('Calibration', 'Adjusting or evaluating intervals to achieve a coverage target.'),
('Quantile or percentile', 'A location within ordered values; 0.9 quantile corresponds to the 90th percentile.'),
('Overfitting', 'Learning training details that do not generalize well.'),
('Data leakage', 'Using information that would be unavailable at prediction time.'),
('Distribution shift', 'The data behavior changes over time or between populations.'),
('Reproduction', 'Reimplementing or rerunning a published experiment with attribution.')], [1.55, 5.1])
h('Sources and where each explanation comes from')
p('[1] Xu, C. and Xie, Y. (2021). Conformal prediction interval for dynamic time-series. ICML, PMLR 139, pages 11559-11569. Read Algorithm 1 and Section 5.1.')
link('Open the original research paper', 'https://proceedings.mlr.press/v139/xu21h.html')
p('[2] Author implementation, hamrel-cxu/EnbPI, main-branch commit 60cd5b7530eb954b02ae94da967111f5b5c3c01b. The project retains source files, license, and hashes.')
link('Open the authors\' code repository', 'https://github.com/hamrel-cxu/EnbPI')
link('[3] scikit-learn 1.6 RandomForestRegressor documentation', 'https://scikit-learn.org/1.6/modules/generated/sklearn.ensemble.RandomForestRegressor.html')
link('[4] NumPy 2.2 percentile documentation', 'https://numpy.org/doc/2.2/reference/generated/numpy.percentile.html')
p('Project evidence: configs/reproduction.json; src/enbpi_project; tests/test_reproduction.py; outputs/solar_reproduction; docs/reproduction-notes.md. All small teaching examples in this guide are invented to illustrate arithmetic; the results tables and charts come from the saved experiment.')

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
print('Native equation paragraphs:', len(doc.element.xpath('//m:oMathPara')))
